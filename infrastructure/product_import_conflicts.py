"""
Read-only database conflict analysis for XLSX product imports (Phase B1).

Stream-batches validated XLSX rows against ProductMaster with bounded
memory.  No database writes.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import date as dt_date, datetime as dt_datetime
from decimal import Decimal
from math import isfinite
from pathlib import Path

from infrastructure.import_constants import MAX_IMPORT_ROWS
from typing import Tuple, Dict, List, Set

import openpyxl
from infrastructure.product_import_preview import ImportColumnMapping
from infrastructure.product_import_identity import (
    ImportSourceSignature, fingerprint_import_source, verify_import_source)


# ── Typed result models ──────────────────────────────────────────────


@dataclass(frozen=True)
class IncomingProduct:
    barcode: str
    name: str
    stock: int
    price: float
    expiry_date: str


@dataclass(frozen=True)
class ExistingProduct:
    barcode: str
    name: str
    stock: int
    price: float
    expiry_date: str


@dataclass(frozen=True)
class ConflictRecord:
    barcode: str
    changed_fields: Tuple[str, ...]


@dataclass(frozen=True)
class ConflictDetail:
    """One field-level difference between incoming and current DB values."""
    barcode: str
    field: str          # English internal name: Name, Stock, Price, ExpiryDate
    current_value: str  # formatted for display
    incoming_value: str  # formatted for display


@dataclass(frozen=True)
class ImportRowErr:
    row_number: int
    barcode: str
    code: str
    message: str


@dataclass(frozen=True)
class ImportConflictResult:
    ok: bool
    cancelled: bool
    error_message: str = ""
    file_name: str = ""
    sheet_name: str = ""
    scanned_rows: int = 0
    valid_rows: int = 0
    invalid_rows: int = 0
    duplicate_barcodes: int = 0
    classified_rows: int = 0
    new_barcodes: int = 0
    unchanged_existing: int = 0
    changed_existing: int = 0
    conflict_samples: Tuple[ConflictRecord, ...] = ()
    conflict_details: Tuple[ConflictDetail, ...] = ()
    errors: Tuple[ImportRowErr, ...] = ()
    sample_rows: Tuple[Tuple[str, str, int, float, str], ...] = ()
    source_signature: ImportSourceSignature | None = None

    @classmethod
    def _make(cls, ok, cancelled, msg, fn, sn, scanned, valid, invalid,
              dupes, classified, new, unchanged, changed,
              conflicts, details, errors, samples, signature=None):
        return cls(
            ok=ok, cancelled=cancelled, error_message=msg,
            file_name=fn, sheet_name=sn,
            scanned_rows=scanned, valid_rows=valid,
            invalid_rows=invalid, duplicate_barcodes=dupes,
            classified_rows=classified,
            new_barcodes=new, unchanged_existing=unchanged,
            changed_existing=changed,
            conflict_samples=tuple(conflicts),
            conflict_details=tuple(details),
            errors=tuple(errors),
            sample_rows=tuple(samples),
            source_signature=signature,
        )

    @classmethod
    def success(cls, fn, sn, scanned, valid, invalid, dupes, classified,
                new, unchanged, changed, conflicts, details, errors, samples,
                signature=None):
        return cls._make(True, False, "", fn, sn, scanned, valid, invalid,
                         dupes, classified, new, unchanged, changed,
                         conflicts, details, errors, samples, signature=signature)

    @classmethod
    def cancelled(cls, fn, sn, scanned, valid, invalid, dupes, classified,
                  new, unchanged, changed, conflicts, details, errors, samples):
        return cls._make(False, True, "Η ανάλυση ακυρώθηκε.",
                         fn, sn, scanned, valid, invalid, dupes,
                         classified, new, unchanged, changed,
                         conflicts, details, errors, samples)

    @classmethod
    def partial(cls, msg, fn, sn, scanned, valid, invalid, dupes,
                classified, new, unchanged, changed,
                conflicts, details, errors, samples):
        return cls._make(False, False, msg, fn, sn, scanned, valid, invalid,
                         dupes, classified, new, unchanged, changed,
                         conflicts, details, errors, samples)

    @classmethod
    def failure(cls, fn, sn, msg):
        return cls._make(False, False, msg, fn, sn, 0, 0, 0, 0, 0, 0, 0, 0,
                         (), (), (), ())


# ── Internal helpers ──────────────────────────────────────────────────


def _connect_ro(db_path: str) -> sqlite3.Connection:
    uri = Path(db_path).absolute().as_uri()
    conn = sqlite3.connect(uri + "?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def _normalize_barcode(raw) -> str | None:
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, str):
        s = raw.strip()
        return s if s else None
    if isinstance(raw, (int, float)):
        if not isfinite(raw):
            return None
        if isinstance(raw, float) and raw != int(raw):
            return None
        s = str(int(raw))
        return s if len(s) <= 15 else None
    return None


def _normalize_name(raw) -> str | None:
    if isinstance(raw, str):
        s = raw.strip()
        return s if s else None
    return None


def _normalize_stock(raw) -> int | None:
    if isinstance(raw, bool) or raw is None:
        return None
    if isinstance(raw, int) and raw >= 0:
        return raw
    if isinstance(raw, float) and isfinite(raw):
        if raw == int(raw):
            s = int(raw)
            return s if s >= 0 else None
    return None


def _normalize_price(raw) -> float | None:
    if isinstance(raw, bool) or raw is None:
        return None
    if isinstance(raw, (int, float)):
        p = float(raw)
        if p >= 0 and isfinite(p):
            return p
    return None


def _normalize_expiry(raw) -> str | None:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return ""
    if isinstance(raw, (dt_date, dt_datetime)):
        return raw.strftime("%Y-%m-%d")
    if isinstance(raw, str):
        s = raw.strip()
        try:
            dt_date.fromisoformat(s)
            return s
        except ValueError:
            return None
    return None


# ── Batched classification ───────────────────────────────────────────


def _fmt_val(val) -> str:
    """Format a DB/incoming value for display.  Blank → '—'."""
    if val is None:
        return "—"
    s = str(val).strip()
    if s == "" or s == "0.0":
        if isinstance(val, float):
            return "—"
    return s if s else "—"


def _classify_batch(batch: List[IncomingProduct], conn: sqlite3.Connection,
                    cancel_event, new_count: int, unchanged_count: int,
                    changed_count: int, conflicts: List[ConflictRecord],
                    details: List[ConflictDetail],
                    MAX_CONFLICT_SAMPLES: int,
                    MAX_CONFLICT_DETAILS: int,
                    ) -> Tuple[int, bool, int, int, int]:
    """Classify a batch. Returns (processed, cancelled, new, unchanged, changed)."""
    barcodes = [p.barcode for p in batch]
    placeholders = ",".join("?" for _ in barcodes)
    rows = conn.execute(
        f"SELECT Barcode, Name, Stock, Price, ExpiryDate "
        f"FROM ProductMaster WHERE Barcode IN ({placeholders})",
        barcodes).fetchall()

    db_data: Dict[str, ExistingProduct] = {}
    for r in rows:
        db_data[r["Barcode"]] = ExistingProduct(
            barcode=r["Barcode"], name=r["Name"] or "",
            stock=r["Stock"] or 0, price=r["Price"] or 0.0,
            expiry_date=r["ExpiryDate"] or "")

    processed = 0
    for prod in batch:
        if cancel_event and cancel_event.is_set():
            return processed, True, new_count, unchanged_count, changed_count
        existing = db_data.get(prod.barcode)
        if existing is None:
            new_count += 1
        else:
            changed_fields: list[str] = []
            # Check each field and build detail records
            if prod.name != existing.name:
                changed_fields.append("Name")
                if len(details) < MAX_CONFLICT_DETAILS:
                    details.append(ConflictDetail(
                        prod.barcode, "Name",
                        _fmt_val(existing.name), _fmt_val(prod.name)))
            if prod.stock != existing.stock:
                changed_fields.append("Stock")
                if len(details) < MAX_CONFLICT_DETAILS:
                    details.append(ConflictDetail(
                        prod.barcode, "Stock",
                        _fmt_val(existing.stock), _fmt_val(prod.stock)))
            if Decimal(str(prod.price)) != Decimal(str(existing.price)):
                changed_fields.append("Price")
                if len(details) < MAX_CONFLICT_DETAILS:
                    details.append(ConflictDetail(
                        prod.barcode, "Price",
                        _fmt_val(existing.price), _fmt_val(prod.price)))
            if prod.expiry_date != existing.expiry_date:
                changed_fields.append("ExpiryDate")
                if len(details) < MAX_CONFLICT_DETAILS:
                    details.append(ConflictDetail(
                        prod.barcode, "ExpiryDate",
                        _fmt_val(existing.expiry_date),
                        _fmt_val(prod.expiry_date or "")))
            if changed_fields:
                changed_count += 1
                if len(conflicts) < MAX_CONFLICT_SAMPLES:
                    conflicts.append(
                        ConflictRecord(prod.barcode, tuple(changed_fields)))
            else:
                unchanged_count += 1
        processed += 1
    return processed, False, new_count, unchanged_count, changed_count


# ── Public API ───────────────────────────────────────────────────────


def analyze_import_conflicts(
    file_path: str,
    mapping: ImportColumnMapping,
    db_path: str,
    sheet_name: str | None = None,
    cancel_event=None,
    max_rows: int = MAX_IMPORT_ROWS,
) -> ImportConflictResult:
    MAX_ERRORS = 200
    MAX_SAMPLES = 20
    MAX_CONFLICT_SAMPLES = 50
    MAX_CONFLICT_DETAILS = 200
    BATCH_SIZE = 500

    # Mapping validation
    map_cols = [
        mapping.barcode_column, mapping.name_column,
        mapping.stock_column, mapping.price_column,
        mapping.expiry_date_column]
    if len(set(map_cols)) != 5:
        return ImportConflictResult.failure(
            file_path, sheet_name or "—",
            "Κάθε στήλη αντιστοίχισης πρέπει να είναι μοναδική.")

    # Open DB read-only before scanning XLSX
    conn = None
    try:
        conn = _connect_ro(db_path)
        info = conn.execute(
            "SELECT name FROM pragma_table_info('ProductMaster')").fetchall()
        existing_cols = {r["name"] for r in info}
        for col in ["Barcode", "Name", "Stock", "Price", "ExpiryDate"]:
            if col not in existing_cols:
                conn.close()
                return ImportConflictResult.failure(
                    file_path, sheet_name or "—",
                    f"Λείπει η στήλη {col} από τον ProductMaster.")
    except sqlite3.Error as e:
        if conn:
            conn.close()
        return ImportConflictResult.failure(
            file_path, sheet_name or "—",
            f"Σφάλμα βάσης δεδομένων: {e}")

    # Create source signature before analysis
    try:
        sig_before = fingerprint_import_source(file_path, mapping)
    except Exception as e:
        if conn:
            conn.close()
        return ImportConflictResult.failure(
            file_path, sheet_name or "—",
            f"Αδυναμία ταυτοποίησης αρχείου: {e}")

    # State
    scanned = 0
    valid = 0
    invalid = 0
    classified = 0
    new_count = 0
    unchanged_count = 0
    changed_count = 0
    errors: List[ImportRowErr] = []
    samples: List[Tuple[str, str, int, float, str]] = []
    conflicts: List[ConflictRecord] = []
    details: List[ConflictDetail] = []
    seen_barcodes: Set[str] = set()
    dupe_barcodes: Set[str] = set()
    batch: List[IncomingProduct] = []
    cancelled = False
    limit_reached = False

    wb = None
    try:
        wb = openpyxl.load_workbook(
            file_path, read_only=True, data_only=True, keep_links=False)
    except Exception as e:
        conn.close()
        return ImportConflictResult.failure(
            file_path, sheet_name or "—",
            f"Αδυναμία ανοίγματος αρχείου: {e}")

    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        active_sheet = ws.title
        col_map: dict[str, int] = {}

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=1, values_only=True), start=1):

            if cancel_event and cancel_event.is_set():
                cancelled = True
                break

            if row_idx == 1:
                headers = tuple(
                    str(c).strip() if c is not None else "" for c in row)
                if len(set(headers)) != len(headers):
                    wb.close()
                    conn.close()
                    return ImportConflictResult.failure(
                        file_path, active_sheet,
                        "Η κεφαλίδα περιέχει διπλότυπα ονόματα στηλών.")
                for key, col_name in [
                    ("barcode", mapping.barcode_column),
                    ("name", mapping.name_column),
                    ("stock", mapping.stock_column),
                    ("price", mapping.price_column),
                    ("expiry", mapping.expiry_date_column)]:
                    try:
                        col_map[key] = headers.index(col_name)
                    except ValueError:
                        wb.close()
                        conn.close()
                        return ImportConflictResult.failure(
                            file_path, active_sheet,
                            f"Στήλη '{col_name}' δεν βρέθηκε.")
                continue

            if scanned >= max_rows:
                limit_reached = True
                if len(errors) < MAX_ERRORS:
                    errors.append(ImportRowErr(
                        row_idx, "", "ROWS",
                        f"Υπέρβαση ορίου {max_rows:,} γραμμών."))
                break

            scanned += 1

            def _col(key):
                idx = col_map[key]
                return row[idx] if idx < len(row) else None

            barcode = _normalize_barcode(_col("barcode"))
            name = _normalize_name(_col("name"))
            stock = _normalize_stock(_col("stock"))
            price = _normalize_price(_col("price"))
            expiry = _normalize_expiry(_col("expiry"))

            row_errs = []
            if barcode is None:
                row_errs.append("Μη έγκυρο barcode.")
            if name is None:
                row_errs.append("Κενό όνομα.")
            if stock is None:
                row_errs.append("Μη έγκυρο απόθεμα.")
            if price is None:
                row_errs.append("Μη έγκυρη τιμή.")
            if expiry is None:
                row_errs.append("Μη έγκυρη ημ. λήξης.")

            if row_errs:
                invalid += 1
                if len(errors) < MAX_ERRORS:
                    errors.append(ImportRowErr(
                        row_idx, barcode or "", "VALIDATION",
                        " · ".join(row_errs)))
                continue

            if barcode in seen_barcodes:
                dupe_barcodes.add(barcode)
                invalid += 1
                if len(errors) < MAX_ERRORS:
                    errors.append(ImportRowErr(
                        row_idx, barcode, "DUPLICATE",
                        f"Διπλότυπο barcode: {barcode}"))
                continue

            seen_barcodes.add(barcode)
            prod = IncomingProduct(
                barcode, name, stock, price, expiry or "")
            batch.append(prod)
            valid += 1
            if len(samples) < MAX_SAMPLES:
                samples.append(
                    (barcode, name, stock, price, expiry or ""))

            # Flush when batch is full
            if len(batch) >= BATCH_SIZE:
                processed, batch_cancelled, new_count, unchanged_count, changed_count = (
                    _classify_batch(batch, conn, cancel_event, new_count,
                                    unchanged_count, changed_count, conflicts,
                                    details,
                                    MAX_CONFLICT_SAMPLES, MAX_CONFLICT_DETAILS))
                classified += processed
                batch.clear()
                if batch_cancelled:
                    cancelled = True
                    break

        # Flush remaining
        if batch and not cancelled:
            processed, batch_cancelled, new_count, unchanged_count, changed_count = (
                _classify_batch(batch, conn, cancel_event, new_count,
                                unchanged_count, changed_count, conflicts,
                                details,
                                MAX_CONFLICT_SAMPLES, MAX_CONFLICT_DETAILS))
            classified += processed
            batch.clear()
            if batch_cancelled:
                cancelled = True
            batch.clear()

    except Exception as e:
        conn.close()
        return ImportConflictResult.failure(
            file_path, sheet_name or "—",
            f"Σφάλμα κατά την ανάγνωση: {e}")
    finally:
        if wb:
            wb.close()

    conn.close()
    conn = None
    dupes = len(dupe_barcodes)

    if cancelled:
        return ImportConflictResult.cancelled(
            file_path, active_sheet, scanned, valid, invalid, dupes,
            classified, new_count, unchanged_count, changed_count,
            conflicts, details, errors, samples)

    if limit_reached:
        return ImportConflictResult.partial(
            f"Υπέρβαση ορίου {max_rows:,} γραμμών.",
            file_path, active_sheet, scanned, valid, invalid, dupes,
            classified, new_count, unchanged_count, changed_count,
            conflicts, details, errors, samples)

    # Verify source signature hasn't changed during analysis
    try:
        if not verify_import_source(sig_before, file_path, mapping):
            return ImportConflictResult.failure(
                file_path, active_sheet,
                "Το αρχείο Excel άλλαξε κατά την ανάλυση. "
                "Επιλέξτε το ξανά.")
    except Exception:
        return ImportConflictResult.failure(
            file_path, active_sheet,
            "Αδυναμία επαλήθευσης ταυτότητας αρχείου.")

    return ImportConflictResult.success(
        file_path, active_sheet, scanned, valid, invalid, dupes,
        classified, new_count, unchanged_count, changed_count,
        conflicts, details, errors, samples, signature=sig_before)
